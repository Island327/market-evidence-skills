"""Read-only, explicitly scoped descriptive analysis; no network or database access."""

import argparse
import csv
import hashlib
import io
import json
import re
import statistics
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

VERSION = "1.0.0"
FIELDS = {
    "record_id", "project_id", "date", "stage", "brand", "region", "model",
    "configuration", "currency", "quantity_unit", "tax_basis", "amount_basis",
    "amount", "quantity", "unit_price",
}
SCOPE_FIELDS = ("stage", "currency", "quantity_unit", "tax_basis", "amount_basis")
DIMENSIONS = {"brand", "region", "model", "stage", "month"}
NUMBER = re.compile(r"[+-]?(?:[0-9]+|[0-9]{1,3}(?:,[0-9]{3})+)(?:\.[0-9]+)?\Z")


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def text_cell(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return str(value).strip()


def number(value, positive=False):
    """No locale guessing, currency stripping, infinity, or blank-to-zero coercion."""
    value = text_cell(value)
    if not NUMBER.fullmatch(value):
        return None
    try:
        parsed = Decimal(value.replace(",", ""))
    except InvalidOperation:
        return None
    if not parsed.is_finite() or parsed < 0 or (positive and parsed == 0):
        return None
    return parsed


def iso_date(value):
    value = text_cell(value)
    if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", value):
        raise ValueError("日期须为 YYYY-MM-DD")
    return date.fromisoformat(value)


def validate_config(config):
    required = {"columns", "scope", "start_date", "end_date", "group_by", "derive_unit_price", "price_relative_tolerance"}
    if not isinstance(config, dict) or set(config) != required:
        raise ValueError("配置键须为: " + ", ".join(sorted(required)))
    columns = config["columns"]
    if not isinstance(columns, dict) or set(columns) != FIELDS:
        raise ValueError("columns 必须映射全部规范字段，见示例配置")
    if any(not isinstance(v, str) or not v.strip() for v in columns.values()):
        raise ValueError("字段映射必须是非空表头")
    if len(set(columns.values())) != len(columns):
        raise ValueError("不同规范字段不能映射到同一列")
    scope = config["scope"]
    if not isinstance(scope, dict) or set(scope) != set(SCOPE_FIELDS):
        raise ValueError("scope 必须声明阶段、币种、数量单位、税口径及金额层级")
    if any(not isinstance(scope[f], str) or not scope[f].strip() for f in SCOPE_FIELDS):
        raise ValueError("每次运行每个范围字段仅接受一个明确值；不同口径请分开运行")
    if scope["amount_basis"] != "line":
        raise ValueError("此工具仅分析已核定的明细行金额；包价/项目总价先另行核实粒度")
    if iso_date(config["start_date"]) > iso_date(config["end_date"]):
        raise ValueError("开始日期晚于结束日期")
    groups = config["group_by"]
    if not isinstance(groups, list) or not groups or any(not isinstance(g, str) or g not in DIMENSIONS for g in groups):
        raise ValueError("group_by 使用 brand/region/model/stage/month 的非空列表")
    if len(set(groups)) != len(groups):
        raise ValueError("group_by 不能重复")
    if not isinstance(config["derive_unit_price"], bool):
        raise ValueError("derive_unit_price 必须为布尔值")
    tolerance = number(config["price_relative_tolerance"])
    if tolerance is None or tolerance > 1:
        raise ValueError("price_relative_tolerance 须在 0 到 1 之间")


def checked_rows(headers, values, config):
    headers = [text_cell(h) for h in headers]
    if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("首行须为非空且唯一的表头；请先明确清理标题、合并表头及空列")
    missing = set(config["columns"].values()) - set(headers)
    if missing:
        raise ValueError("缺少表头: " + ", ".join(sorted(missing)))
    records = []
    for row_number, cells in values:
        if all(text_cell(c) == "" for c in cells):
            continue
        if len(cells) != len(headers):
            raise ValueError(f"第 {row_number} 行列数与表头不同")
        raw = dict(zip(headers, cells))
        record = {key: text_cell(raw[value]) for key, value in config["columns"].items()}
        record["_row"] = row_number
        records.append(record)
    return records


def read_table(path, config, sheet=None):
    if path.suffix.lower() == ".csv":
        if sheet:
            raise ValueError("CSV 不接受 --sheet")
        rows = csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig"), newline=""))
        header = next(rows, None)
        if header is None:
            raise ValueError("输入为空文件")
        return checked_rows(header, enumerate(rows, start=2), config)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("仅支持 UTF-8 CSV 或 XLSX；请先转换其他格式并保留原件")
    if not sheet:
        raise ValueError("XLSX 必须使用 --sheet 指定明细工作表")
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("读取 XLSX 需要 openpyxl，参见可选依赖说明") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError("找不到指定工作表")
        worksheet = workbook[sheet]
        worksheet.reset_dimensions()  # Read actual cells even when producer metadata is wrong.
        rows = worksheet.iter_rows()
        first = next(rows, None)
        if first is None:
            raise ValueError("工作表为空")
        headers = [text_cell(c.value) for c in first]
        needed = set(config["columns"].values())
        indices = {i for i, name in enumerate(headers) if name in needed}
        def values():
            for row_number, cells in enumerate(rows, start=2):
                if any(i < len(cells) and cells[i].data_type in {"f", "e"} for i in indices):
                    raise ValueError(f"第 {row_number} 行含分析字段公式或 Excel 错误，请核实重算并提供值副本")
                raw = [c.value for c in cells]
                if len(raw) < len(headers):
                    raw += [None] * (len(headers) - len(raw))
                yield row_number, raw
        return checked_rows(headers, values(), config)
    finally:
        workbook.close()


def metrics(rows):
    amounts = [r["_amount"] for r in rows if r["_amount"] is not None]
    quantities = [r["_quantity"] for r in rows if r["_quantity"] is not None]
    return {
        "rows": len(rows),
        "amount_n": len(amounts), "amount_total": sum(amounts, Decimal(0)) if amounts else None,
        "quantity_n": len(quantities), "quantity_total": sum(quantities, Decimal(0)) if quantities else None,
        "project_count": len({r["project_id"] for r in rows if r["project_id"]}),
        "project_missing_rows": sum(not r["project_id"] for r in rows),
    }


def analyze(records, config):
    validate_config(config)
    review, eligible = [], []
    ids = defaultdict(list)
    for record in records:
        if record["record_id"]:
            ids[record["record_id"]].append(record)
    conflicts = {key for key, rows in ids.items() if len({tuple(r[f] for f in sorted(FIELDS)) for r in rows}) > 1}
    seen = set()
    start, end = iso_date(config["start_date"]), iso_date(config["end_date"])
    tolerance = number(config["price_relative_tolerance"])

    def issue(row, impact, reason, metric=""):
        review.append({"source_row": row["_row"], "record_id": row["record_id"], "impact": impact, "metric": metric, "reason": reason})

    for original in records:
        row = dict(original)
        rid = row["record_id"]
        if not rid:
            issue(row, "row_excluded", "missing_record_id")
            continue
        if rid in conflicts:
            issue(row, "row_excluded", "conflicting_record_id")
            continue
        if rid in seen:
            issue(row, "row_excluded", "identical_record_id_duplicate")
            continue
        seen.add(rid)
        outside = [f for f in SCOPE_FIELDS if row[f] != config["scope"][f]]
        try:
            day = iso_date(row["date"])
            if not start <= day <= end:
                outside.append("date_range")
        except ValueError:
            outside.append("invalid_date")
        if outside:
            issue(row, "row_excluded", "outside_scope:" + ",".join(outside))
            continue
        row["month"] = row["date"][:7]
        row["_amount"] = number(row["amount"])
        row["_quantity"] = number(row["quantity"], positive=True)
        row["_price"] = number(row["unit_price"], positive=True)
        for field in ("amount", "quantity"):
            if row["_" + field] is None:
                issue(row, "metric_excluded", "missing_or_invalid_" + field, field)
        if not row["project_id"]:
            issue(row, "metric_excluded", "missing_project_id", "project_count")
        for field in config["group_by"]:
            if not row[field]:
                issue(row, "warning", "missing_group_dimension:" + field)
        can_derive = row["_amount"] is not None and row["_quantity"] is not None
        if row["_price"] is not None and can_derive:
            expected = row["_price"] * row["_quantity"]
            if abs(row["_amount"] - expected) > tolerance * max(row["_amount"], expected):
                row["_price"] = None
                issue(row, "metric_excluded", "unit_price_amount_conflict", "price")
        elif not row["unit_price"] and can_derive and config["derive_unit_price"] and row["_amount"] > 0:
            row["_price"] = row["_amount"] / row["_quantity"]
            issue(row, "warning", "unit_price_derived_from_line_amount", "price")
        if row["_price"] is None:
            issue(row, "metric_excluded", "missing_or_invalid_unit_price", "price")
        elif not row["model"] or not row["configuration"]:
            row["_price"] = None
            issue(row, "metric_excluded", "missing_price_comparability_key", "price")
        eligible.append(row)

    total = metrics(eligible)
    grouped = defaultdict(list)
    price_groups = defaultdict(list)
    for row in eligible:
        # Empty is an explicit unknown bucket; do not collide with a real brand named "Unknown".
        grouped[tuple(row[field] for field in config["group_by"])].append(row)
        if row["_price"] is not None:
            price_groups[(row["model"], row["configuration"], row["month"])].append(row["_price"])
    groups = []
    for keys, rows in sorted(grouped.items()):
        result = dict(zip(config["group_by"], keys))
        result.update(metrics(rows))
        for metric in ("amount", "quantity"):
            denominator = total[metric + "_total"]
            numerator = result[metric + "_total"]
            result[metric + "_share"] = numerator / denominator if numerator is not None and denominator else None
        groups.append(result)
    prices = [
        {"model": key[0], "configuration": key[1], "month": key[2], "n": len(values),
         "median_unit_price": statistics.median(values), "min_unit_price": min(values), "max_unit_price": max(values)}
        for key, values in sorted(price_groups.items())
    ]
    row_exclusions = {r["source_row"] for r in review if r["impact"] == "row_excluded"}
    summary = {
        "method_version": VERSION,
        "scope": config["scope"], "period": {"start": config["start_date"], "end": config["end_date"]},
        "input_rows": len(records), "eligible_rows": len(eligible), "excluded_rows": len(row_exclusions),
        "metrics": total, "price_n": sum(p["n"] for p in prices),
        "review_issue_counts": dict(sorted(Counter(r["reason"] for r in review).items())),
        "groups": groups, "price_strata": prices,
        "limitations": [
            "Descriptive observed sample only; not total market size or causal evidence.",
            "Project counts are non-additive across groups; blank dimension values are retained as an unknown bucket.",
            "Amount and quantity have separate valid populations; do not divide the two overall totals to infer a unit price.",
            "Price strata fix model/configuration/month plus the declared currency/unit/tax/stage; commercial comparability still needs review.",
            "Distinct record IDs are not evidence of distinct business events; cross-source/lifecycle deduplication is outside this helper.",
            "Shares exclude invalid metric values; a null group metric is unknown, not confirmed zero.",
        ],
    }
    return summary, review


def json_bytes(data):
    def convert(value):
        if isinstance(value, Decimal):
            return format(value, "f")  # Exact decimal strings; never binary float money.
        raise TypeError(type(value).__name__)
    return (json.dumps(data, ensure_ascii=False, indent=2, default=convert) + "\n").encode("utf-8")


def csv_value(value):
    if value is None:
        return ""
    if isinstance(value, str) and value and (value[:1] in "\t\r\n" or value.lstrip().startswith(("=", "+", "-", "@"))):
        return "'" + value
    return value


def csv_bytes(rows, fields):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields)
    writer.writeheader()
    for row in rows:
        writer.writerow({key: csv_value(row.get(key)) for key in fields})
    return stream.getvalue().encode("utf-8-sig")


def report(summary):
    m = summary["metrics"]
    show = lambda v: "未获得有效值" if v is None else str(v)
    return (
        "# 观察样本基础分析\n\n"
        f"输入 {summary['input_rows']} 行；范围内且通过身份检查 {summary['eligible_rows']} 行；排除 {summary['excluded_rows']} 行。\n\n"
        f"- 金额合计：{show(m['amount_total'])}；有效记录 {m['amount_n']}。\n"
        f"- 数量合计：{show(m['quantity_total'])}；有效记录 {m['quantity_n']}。\n"
        f"- 去重项目数：{m['project_count']}；缺项目键 {m['project_missing_rows']} 行。\n"
        f"- 可用于已声明价格分层的记录：{summary['price_n']}。\n\n"
        "范围、币种、税口径、数量单位和日期见 summary.json；分组见 groups.csv；价格见 price_strata.csv；问题见 review.csv。\n\n"
        "这份文件只报告基础统计。金额与数量的有效母集可能不同，不能直接以两个总计相除计算均价。空值不等于零；分组项目数不能相加；样本份额不等于全市场份额。价格分层不自动证明商业可比性。\n\n"
        "下一步：按任务核对原文、异常值及样本覆盖，结合参考方法撰写发现、解释和建议。\n"
    ).encode("utf-8")


def run(input_path, config_path, output_path, sheet=None):
    if output_path.exists():
        raise ValueError("输出目录已存在，请为本次运行选择新的目录")
    input_digest = sha256(input_path.read_bytes())
    config_data = config_path.read_bytes()
    config = json.loads(config_data)
    validate_config(config)
    records = read_table(input_path, config, sheet)
    if sha256(input_path.read_bytes()) != input_digest:
        raise ValueError("读取期间输入文件发生变化，请重新冻结输入")
    summary, review = analyze(records, config)
    group_fields = config["group_by"] + list(metrics([])) + ["amount_share", "quantity_share"]
    artifacts = {
        "summary.json": json_bytes(summary),
        "groups.csv": csv_bytes(summary["groups"], group_fields),
        "price_strata.csv": csv_bytes(summary["price_strata"], ["model", "configuration", "month", "n", "median_unit_price", "min_unit_price", "max_unit_price"]),
        "review.csv": csv_bytes(review, ["source_row", "record_id", "impact", "metric", "reason"]),
        "report.md": report(summary),
    }
    manifest = {
        "schema_version": "1.0", "method_version": VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input": {"name": input_path.name, "sha256": input_digest, "sheet": sheet, "nonempty_rows": len(records)},
        "config": {"name": config_path.name, "sha256": sha256(config_data), "effective": config},
        "script_sha256": sha256(Path(__file__).read_bytes()),
        "outputs": [{"name": name, "sha256": sha256(data)} for name, data in artifacts.items()],
        "state": "generated", "sent": False, "recipient_import_verified": False,
    }
    output_path.mkdir(parents=True, exist_ok=False)
    for name, data in artifacts.items():
        (output_path / name).write_bytes(data)
    (output_path / "manifest.json").write_bytes(json_bytes(manifest))
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--sheet")
    args = parser.parse_args()
    try:
        summary = run(args.input, args.config, args.out, args.sheet)
    except (ValueError, OSError, csv.Error) as exc:
        print("分析未完成: " + str(exc), file=sys.stderr)
        return 2
    print(f"已生成 6 个输出文件；有效 {summary['eligible_rows']} 行，排除 {summary['excluded_rows']} 行。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
